import io
from fastapi.responses import StreamingResponse
from services.llm_integration.llm_interface import get_llm
from .generate_material_utils import GenerateMaterialRequest, GenerateMaterialResponse, generate_material_prompt, MaterialType

# DOCX generation
from docx import Document
from markdown2 import markdown
from bs4 import BeautifulSoup

def generate_docx(markdown_text: str) -> io.BytesIO:
    html = markdown(markdown_text)
    soup = BeautifulSoup(html, "html.parser")
    doc = Document()

    for element in soup.recursiveChildGenerator():
        if element.name == 'h1':
            doc.add_heading(element.get_text(), level=1)
        elif element.name == 'h2':
            doc.add_heading(element.get_text(), level=2)
        elif element.name == 'h3':
            doc.add_heading(element.get_text(), level=3)
        elif element.name == 'p':
            doc.add_paragraph(element.get_text())
        elif element.name == 'ul':
            for li in element.find_all("li"):
                doc.add_paragraph(li.get_text(), style="List Bullet")
        elif element.name == 'ol':
            for li in element.find_all("li"):
                doc.add_paragraph(li.get_text(), style="List Number")

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

# PDF generation
from reportlab.platypus import SimpleDocTemplate, Paragraph
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.pagesizes import A4

def generate_pdf(markdown_content: str) -> io.BytesIO:
    html = markdown(markdown_content)
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4)
    styles = getSampleStyleSheet()
    story = [Paragraph(html, styles["Normal"])]
    doc.build(story)
    buffer.seek(0)
    return buffer

# Main material handler
def material(request: GenerateMaterialRequest):
    llm = get_llm(request.model)

    try:
        response: GenerateMaterialResponse = llm.generate_text(
            prompt=generate_material_prompt(request),
            response_model=GenerateMaterialResponse
        )
        generated_string = response.material
        filename_base = request.title.replace(" ", "_").lower()

        fmt = request.type_of_file
        ext = fmt.value.lower()

        mime_types = {
            MaterialType.MARKDOWN: "text/markdown",
            MaterialType.PDF: "application/pdf",
            MaterialType.DOCX: "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        }

        media_type = mime_types[fmt]
        filename = f"{filename_base}.{ext}"

        if fmt == MaterialType.MARKDOWN:
            file_bytes = io.BytesIO(generated_string.encode("utf-8"))
        elif fmt == MaterialType.DOCX:
            file_bytes = generate_docx(generated_string)
        elif fmt == MaterialType.PDF:
            file_bytes = generate_pdf(generated_string)
        else:
            raise ValueError("Unsupported file format requested.")

        return StreamingResponse(
            file_bytes,
            media_type=media_type,
            headers={"Content-Disposition": f'attachment; filename="{filename}"'}
        )

    except Exception as e:
        raise e
